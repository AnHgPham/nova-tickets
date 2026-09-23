#!/usr/bin/env python3
"""Tạo tài liệu giải pháp kỹ thuật NOVA Tickets theo form HDSD, không logo/tên Viettel."""

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor, Twips

OUT = (
    "/home/loogix/Downloads/NOVA-Tickets-final-source-b3da9d39/"
    "docs/NOVA-Tickets-Tai-lieu-giai-phap-ky-thuat.docx"
)
RED = "EE0033"
NAVY = "1F4D78"


def shade(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, kind="single", sz="4", color="000000"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), kind)
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def font_run(run, name="Times New Roman", size=12, bold=False, italic=False, color="000000"):
    run.font.name = name
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def p_format(p, before=0, after=10, line=1.15, align="justify", first=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    align_map = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    pf.alignment = align_map[align]
    if first is not None:
        pf.first_line_indent = Cm(first)


def add_text(p, text, **kwargs):
    run = p.add_run(text)
    font_run(run, **kwargs)
    return run


def body(doc, text, **kwargs):
    p = doc.add_paragraph()
    p_format(p, before=kwargs.pop("before", 0), after=kwargs.pop("after", 10),
             line=kwargs.pop("line", 1.15), align=kwargs.pop("align", "justify"))
    add_text(p, text, **kwargs)
    return p


def heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    size = {1: 16, 2: 14, 3: 13}[level]
    before = {1: 18, 2: 14, 3: 6}[level]
    after = {1: 18, 2: 6, 3: 6}[level]
    line = {1: 1.15, 2: 1.0, 3: 1.5}[level]
    for run in p.runs:
        font_run(run, size=size, bold=True, color="000000")
    p_format(p, before=before, after=after, line=line, align="left")
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p_format(p, before=0, after=0, line=1.0, align="left")
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def set_table_widths(table, widths_dxa):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths_dxa)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for w in widths_dxa:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[i]))
            tcW.set(qn("w:type"), "dxa")


def fill_header_row(row, labels, fill, border="single"):
    for cell, text in zip(row.cells, labels):
        cell.text = ""
        p = cell.paragraphs[0]
        p_format(p, before=2, after=2, line=1.0, align="left")
        add_text(p, text, size=11, bold=True)
        shade(cell, fill)
        set_cell_border(cell, kind=border)
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        vAlign = tcPr.find(qn("w:vAlign"))
        if vAlign is None:
            vAlign = OxmlElement("w:vAlign")
            tcPr.append(vAlign)
        vAlign.set(qn("w:val"), "center")


def fill_body_cell(cell, text, border="single", bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    p_format(p, before=2, after=2, line=1.0, align="left")
    add_text(p, text, size=size, bold=bold)
    set_cell_border(cell, kind=border)


def add_table(doc, headers, rows, widths, header_fill="FFCC99", border="single"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    fill_header_row(table.rows[0], headers, header_fill, border)
    for r_i, row in enumerate(rows):
        for c_i, value in enumerate(row):
            fill_body_cell(table.rows[r_i + 1].cells[c_i], value, border)
    set_table_widths(table, widths)
    doc.add_paragraph()
    return table


def code_block(doc, text):
    for line in text.strip("\n").split("\n"):
        p = doc.add_paragraph()
        p_format(p, before=0, after=0, line=1.0, align="left")
        add_text(p, line if line else " ", name="Consolas", size=10, color="000000")
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        pPr.append(shd)


def add_footer_header(section):
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False

    hp = section.header.paragraphs[0]
    hp.clear()
    p_format(hp, before=0, after=0, line=1.0, align="right")
    add_text(hp, "NOVA Tickets  ·  Tài liệu giải pháp kỹ thuật", size=10, color=NAVY)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), RED)
    pBdr.append(bottom)
    pPr.append(pBdr)

    for p in section.first_page_header.paragraphs:
        p.text = ""
    for p in section.first_page_footer.paragraphs:
        p.text = ""

    footer = section.footer
    footer.paragraphs[0].text = ""
    table = footer.add_table(rows=1, cols=2, width=Cm(16.2))
    left, right = table.rows[0].cells
    left.text = ""
    lp = left.paragraphs[0]
    p_format(lp, before=0, after=0, line=1.0, align="left")
    add_text(
        lp,
        "Tài liệu chứa nội dung quan trọng. Không chia sẻ trái phép dưới mọi hình thức.",
        size=10,
    )
    right.text = ""
    rp = right.paragraphs[0]
    p_format(rp, before=0, after=0, line=1.0, align="center")

    def fld(paragraph, instr):
        r1 = paragraph.add_run()
        fc1 = OxmlElement("w:fldChar")
        fc1.set(qn("w:fldCharType"), "begin")
        r1._r.append(fc1)
        r2 = paragraph.add_run()
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = f" {instr} "
        r2._r.append(it)
        font_run(r2, size=8, bold=True)
        r3 = paragraph.add_run()
        fc2 = OxmlElement("w:fldChar")
        fc2.set(qn("w:fldCharType"), "end")
        r3._r.append(fc2)

    fld(rp, "PAGE")
    add_text(rp, "/", size=8, bold=True)
    fld(rp, "NUMPAGES")
    for cell in (left, right):
        set_cell_border(cell, kind="nil", sz="0", color="auto")
    set_table_widths(table, [8200, 1600])


def add_toc(doc):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 1"]
    p_format(p, before=18, after=6, line=1.15, align="center")
    add_text(p, "MỤC LỤC", size=16, bold=True)

    note = doc.add_paragraph()
    p_format(note, before=0, after=10, line=1.0, align="left")
    add_text(
        note,
        '(Nhấp chuột phải vào mục lục và chọn "Update Field" để cập nhật số trang sau khi chỉnh sửa tài liệu)',
        size=10,
        italic=True,
    )

    toc = doc.add_paragraph()
    p_format(toc, before=0, after=6, line=1.15, align="left")
    r1 = toc.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    b.set(qn("w:dirty"), "true")
    r1._r.append(b)
    r2 = toc.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = ' TOC \\o "1-3" \\h \\z \\u '
    r2._r.append(it)
    font_run(r2, size=12)
    r3 = toc.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3._r.append(sep)
    add_text(toc, "Mở file bằng Microsoft Word và cập nhật field để hiện số trang.", size=12, italic=True)
    r4 = toc.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r4._r.append(end)
    page_break(doc)


def style_base(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor.from_string("00000A")
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(10)
    pf.line_spacing = 1.15
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), "Times New Roman")


def enable_update_fields(doc):
    settings = doc.settings.element
    uf = settings.find(qn("w:updateFields"))
    if uf is None:
        uf = OxmlElement("w:updateFields")
        settings.insert(0, uf)
    uf.set(qn("w:val"), "true")


def build():
    doc = Document()
    style_base(doc)
    sec = doc.sections[0]
    sec.page_width = Emu(7772400)
    sec.page_height = Emu(10058400)
    sec.left_margin = Emu(914400)
    sec.right_margin = Emu(731520)
    sec.top_margin = Emu(822960)
    sec.bottom_margin = Emu(722630)
    sec.header_distance = Twips(288)
    sec.footer_distance = Twips(0)
    add_footer_header(sec)

    # Cover — text only, no logo
    kicker = doc.add_paragraph()
    p_format(kicker, before=36, after=6, line=1.0, align="center")
    add_text(kicker, "ĐỒ ÁN / BÁO CÁO HỌC PHẦN", size=12, bold=True, color=NAVY)

    title = doc.add_paragraph()
    p_format(title, before=28, after=4, line=1.0, align="center")
    add_text(title, "NOVA TICKETS", size=28, bold=True, color=RED)

    sub = doc.add_paragraph()
    p_format(sub, before=4, after=2, line=1.0, align="center")
    add_text(sub, "Hệ thống web bán vé sự kiện âm nhạc", size=14, bold=True)

    kind = doc.add_paragraph()
    p_format(kind, before=10, after=2, line=1.0, align="center")
    add_text(kind, "Tài liệu giải pháp kỹ thuật", size=16, bold=True)

    meta = doc.add_paragraph()
    p_format(meta, before=8, after=16, line=1.15, align="center")
    add_text(meta, "Phiên bản: 1.0\nNgày cập nhật: 23/09/2026", size=12)

    info_headers = ["Hạng mục", "Nội dung"]
    info_rows = [
        ["Học phần", ""],
        ["Giảng viên", ""],
        ["Sinh viên thực hiện", ""],
        ["MSSV", ""],
        ["Lớp", ""],
        ["Phiên bản tài liệu", "1.0"],
        ["Ngày lập", "23/09/2026"],
    ]
    add_table(doc, info_headers, info_rows, [3200, 6400], header_fill="F2F2F2")
    note = doc.add_paragraph()
    p_format(note, before=0, after=0, line=1.0, align="center")
    add_text(note, "Các ô học phần, giảng viên, sinh viên, MSSV và lớp để trống để điền khi nộp.", size=10, italic=True)
    page_break(doc)

    # Changelog
    ch = doc.add_paragraph()
    p_format(ch, before=18, after=6, line=1.0, align="center")
    add_text(ch, "BẢNG GHI NHẬN THAY ĐỔI TÀI LIỆU", size=12, bold=True)
    legend = doc.add_paragraph()
    p_format(legend, before=0, after=6, line=1.0, align="left")
    add_text(legend, "*A – Tạo mới, M – Sửa đổi, D – Xóa bỏ", size=12)
    add_table(
        doc,
        ["Ngày thay đổi", "Vị trí thay đổi", "Nguồn gốc", "Phiên bản cũ", "Mô tả thay đổi", "Phiên bản mới"],
        [["23/09/2026", "Toàn bộ", "A", "—", "Tạo mới tài liệu giải pháp kỹ thuật NOVA Tickets theo mã nguồn và đặc tả.", "1.0"]],
        [1400, 1600, 1100, 1400, 2800, 1300],
        header_fill="BFBFBF",
        border="dotted",
    )
    page_break(doc)
    add_toc(doc)

    heading(doc, "1. Tổng quan giải pháp", 1)
    heading(doc, "1.1. Mô tả chung", 2)
    body(
        doc,
        "NOVA Tickets là hệ thống web bán vé sự kiện âm nhạc theo sơ đồ ghế. Người dùng xem các đêm diễn đang mở bán, chọn ghế, giữ ghế tạm thời trong 7 phút, xác nhận đặt vé và theo dõi đơn của chính mình. Hệ thống tách storefront khách hàng, console quản trị và REST API backend.",
    )
    body(
        doc,
        "Backend dùng ASP.NET Core 8 Web API, Entity Framework Core và MySQL/TiDB. Frontend dùng React và TypeScript, được biên dịch vào wwwroot và do chính ASP.NET Core phục vụ cùng origin. Trình duyệt gọi API theo một domain; backend kiểm soát cookie xác thực và toàn bộ nghiệp vụ nhạy cảm.",
    )
    body(
        doc,
        "Trạng thái ghế chỉ được quyết định tại backend và cơ sở dữ liệu. Giao diện chỉ hiển thị dữ liệu đã trả về, không được tự xác nhận rằng một ghế còn bán được.",
    )

    heading(doc, "1.2. Mục tiêu và phạm vi", 2)
    add_table(
        doc,
        ["Nhóm", "Mục tiêu / phạm vi"],
        [
            ["Khách hàng", "Xem sự kiện, chọn ghế trực quan, giữ ghế, xác nhận booking và xem đơn vé cá nhân."],
            ["Quản trị", "CRUD sự kiện, địa điểm, suất diễn, khu vực ghế, ghế, loại vé, giá theo suất diễn, người dùng và booking."],
            ["Toàn vẹn dữ liệu", "Không phát hành hai vé cho cùng một ghế trong cùng một suất diễn."],
            ["Bảo mật", "JWT cookie, phân quyền theo vai trò, validation, rate limiting, security headers, audit log và correlation ID."],
            ["Vận hành", "Tự dọn hold quá hạn, chuẩn hóa lỗi API, kịch bản kiểm thử tự động và tài liệu thao tác."],
            ["Thông báo", "Hàng đợi email xác nhận (outbox) để xem trước và gửi lại, chưa nối SMTP thật."],
        ],
        [2800, 6800],
    )

    heading(doc, "1.3. Các vai trò sử dụng", 2)
    add_table(
        doc,
        ["Vai trò", "Chức năng chính", "Giới hạn"],
        [
            ["Khách vãng lai", "Xem nội dung công khai và danh sách sự kiện.", "Không giữ ghế, không xác nhận booking, không vào quản trị."],
            ["Customer", "Đăng ký, đăng nhập, giữ ghế, tạo booking, xem đơn của mình.", "Không đọc hay sửa dữ liệu quản trị."],
            ["Staff", "Hỗ trợ vận hành trong phạm vi quyền được cấp, gồm xem outbox email.", "Không mặc định có toàn quyền quản trị."],
            ["Admin", "Quản lý catalog, suất diễn, sơ đồ ghế, loại vé, giá, người dùng và booking.", "Hủy hoặc xóa phải tuân theo lịch sử giao dịch."],
            ["Scheduled job", "Gọi API dọn hold hết hạn mỗi phút.", "Chỉ dùng endpoint bảo trì đã xác thực bằng khóa riêng."],
        ],
        [2200, 4200, 3200],
    )

    heading(doc, "1.4. Bảng chức năng", 2)
    add_table(
        doc,
        ["Chức năng", "Mô tả", "Đối tượng sử dụng"],
        [
            ["Xem sự kiện và sơ đồ ghế", "Đọc catalog công khai và trạng thái Available, Held, Sold, Blocked.", "Khách vãng lai, Customer"],
            ["Đăng ký / đăng nhập", "Tạo tài khoản và nhận JWT trong cookie HttpOnly, SameSite=Strict.", "Customer"],
            ["Giữ ghế", "Giữ tối đa 8 ghế trong 7 phút bằng hold token.", "Customer"],
            ["Xác nhận booking", "Đổi hold hợp lệ thành booking, ticket và payment nội bộ, có idempotency.", "Customer"],
            ["Đơn của tôi", "Chỉ trả booking thuộc phiên đang đăng nhập.", "Customer"],
            ["Quản trị catalog", "CRUD sự kiện, địa điểm, suất diễn, ghế, loại vé và giá.", "Admin"],
            ["Quản trị người dùng và đơn", "Vô hiệu hóa user thay vì xóa cứng; đổi trạng thái đơn theo luật.", "Admin"],
            ["Email xác nhận", "Lưu HTML/text vào outbox, xem trước và retry khi Failed.", "Customer (xem của mình), Staff, Admin"],
            ["Dọn hold hết hạn", "Trả ghế Held quá hạn về Available.", "Tiến trình có khóa scheduler"],
        ],
        [2800, 4300, 2500],
    )

    heading(doc, "2. API xác thực và phân quyền", 1)
    heading(doc, "2.1. Mô tả chung", 2)
    body(
        doc,
        "Nhóm API xác thực cung cấp đăng ký, đăng nhập, đăng xuất và đọc phiên. Sau khi đăng nhập thành công, backend phát JWT trong cookie HttpOnly và SameSite=Strict. Mỗi request cần xác thực đều được kiểm tra JWT. Vai trò Customer, Staff và Admin giới hạn quyền truy cập.",
    )
    heading(doc, "2.2. Mô tả API", 2)
    add_table(
        doc,
        ["STT", "Tên API", "URL API", "Tham số", "Response / ghi chú"],
        [
            ["1", "Register", "POST /api/auth/register", "email, password, fullName, phone", "201/200 khi tạo user; 400 nếu input sai hoặc email đã tồn tại."],
            ["2", "Login", "POST /api/auth/login", "email, password", "200 và phát cookie JWT; 401 khi sai thông tin hoặc tài khoản không hợp lệ."],
            ["3", "Logout", "POST /api/auth/logout", "Không body", "200; xóa cookie phiên."],
            ["4", "Me", "GET /api/auth/me", "Cookie JWT", "200 trả hồ sơ; 401 nếu chưa xác thực."],
        ],
        [700, 1400, 2500, 2300, 2700],
    )
    heading(doc, "2.3. Mô tả chi tiết xử lý", 2)
    for step in [
        "Client gửi đăng ký hoặc đăng nhập qua HTTPS cùng origin.",
        "Backend validate email, mật khẩu, tên và các trường bắt buộc.",
        "Khi đăng nhập, backend kiểm tra password hash và trạng thái IsActive.",
        "Nếu hợp lệ, backend tạo JWT chứa định danh và role, rồi ghi vào cookie bảo mật.",
        "Controller nghiệp vụ chặn request chưa xác thực bằng 401 và request không đủ quyền bằng 403.",
        "Request có thể mang X-Correlation-ID. Thao tác nhạy cảm được ghi audit để truy vết.",
    ]:
        body(doc, step, before=0, after=4)

    heading(doc, "3. API giữ ghế và xác nhận booking", 1)
    heading(doc, "3.1. Mô tả chung", 2)
    body(
        doc,
        "Đây là nhóm API trọng tâm. Hệ thống không cho tạo booking trực tiếp. Chuỗi bắt buộc là đọc sơ đồ ghế, giữ ghế, rồi xác nhận booking. Cách này xử lý hai khách chọn cùng một ghế, request bị gửi lại vì lỗi mạng, và trường hợp khách bỏ dở.",
    )
    heading(doc, "3.2. Mô tả API", 2)
    add_table(
        doc,
        ["STT", "Tên API", "URL API", "Tham số", "Response / ghi chú"],
        [
            ["1", "Get seat map", "GET /api/performances/{id}/seats", "performanceId", "200; ghế Available, Held, Sold hoặc Blocked."],
            ["2", "Hold seats", "POST /api/seat-holds", "performanceId, seatIds[]", "200 kèm holdToken, expiresAtUtc và tổng tiền. 409 seat_unavailable nếu ghế vừa bị giữ hoặc mua."],
            ["3", "Release hold", "DELETE /api/seat-holds/{holdToken}", "holdToken", "Giải phóng sớm các ghế của token hợp lệ."],
            ["4", "Confirm booking", "POST /api/bookings", "holdToken, thông tin khách, idempotencyKey", "200 tạo booking, ticket và payment. 409 nếu hold hết hạn hoặc không thuộc user."],
            ["5", "My bookings", "GET /api/bookings/mine", "Cookie JWT", "200; chỉ đơn của user hiện tại."],
        ],
        [700, 1600, 2800, 2200, 2300],
    )
    heading(doc, "3.3. Luồng giữ ghế", 2)
    for step in [
        "Khách chọn tối đa 8 ghế đang Available.",
        "Backend kiểm tra user đã đăng nhập, suất diễn OnSale, cửa sổ bán còn hiệu lực và danh sách ghế hợp lệ.",
        "Backend mở transaction ngắn và cập nhật có điều kiện từng hàng SeatInventory: chỉ chuyển sang Held khi đang Available hoặc hold cũ đã quá hạn.",
        "Nếu một ghế không cập nhật được, toàn bộ transaction rollback và API trả HTTP 409 với mã seat_unavailable.",
        "Nếu thành công, hệ thống ghi HoldToken, HeldByUserId, HoldExpiresAtUtc và trả hạn 7 phút cho frontend.",
    ]:
        body(doc, step, before=0, after=4)
    heading(doc, "3.4. Luồng xác nhận booking", 2)
    for step in [
        "Client gửi holdToken, thông tin liên hệ và idempotencyKey.",
        "Nếu đã có booking cùng UserId và IdempotencyKey, backend trả kết quả cũ và không tạo bản ghi mới.",
        "Trong transaction ReadCommitted, backend xác nhận token thuộc user hiện tại, hold chưa hết hạn và mọi ghế vẫn Held.",
        "Backend tạo Booking, từng Ticket và một Payment nội bộ, rồi chuyển các ghế từ Held sang Sold trong cùng một commit.",
        "Nếu ràng buộc hoặc trạng thái không còn hợp lệ, transaction rollback. Frontend không được tự đánh dấu vé thành công.",
    ]:
        body(doc, step, before=0, after=4)

    heading(doc, "3.5. Hai người cùng chọn một ghế", 2)
    body(
        doc,
        "Hai client có thể cùng thấy một ghế là trống vì giao diện cập nhật trễ. Tại thời điểm ghi, chỉ request đầu tiên cập nhật được đúng một hàng. Request còn lại nhận 409. Ràng buộc duy nhất (PerformanceId, SeatId) trên bảng Ticket là lớp phòng vệ cuối: một ghế không có hai vé trong cùng suất diễn.",
    )
    code_block(
        doc,
        """Khách A                         Backend / Database                    Khách B
|-- POST hold(A01) -------------->|
|                                 | UPDATE A01 WHERE Status=Available
|                                 | 1 row affected -> A01 = Held
|<-- 200 + holdToken -------------|
|                                 |<---------------- POST hold(A01) --|
|                                 | UPDATE A01 WHERE Status=Available
|                                 | 0 rows affected
|                                 |----------------> 409 seat_unavailable""",
    )
    body(
        doc,
        "Tiêu chí đạt bắt buộc: hai request giữ cùng một ghế gửi đồng thời phải ra một HTTP 200 và một HTTP 409. Hai HTTP 200 là lỗi nghiêm trọng vì có nguy cơ bán trùng.",
        before=8,
    )

    heading(doc, "4. CRUD quản trị sự kiện và vé", 1)
    heading(doc, "4.1. Mô tả chung", 2)
    body(
        doc,
        "Admin console và Swagger cùng origin cho quản trị viên thao tác dữ liệu vận hành. Thiết kế ưu tiên giữ lịch sử giao dịch. Khi đã có dữ liệu phụ thuộc, hệ thống vô hiệu hóa hoặc chặn thay đổi không an toàn thay vì xóa cứng.",
    )
    heading(doc, "4.2. Danh sách API", 2)
    add_table(
        doc,
        ["Resource", "API", "Ghi chú nghiệp vụ"],
        [
            ["Event", "GET/POST /api/events; GET/PUT/DELETE /api/events/{id}", "Sự kiện Draft không được mở bán."],
            ["Venue", "GET/POST /api/venues; GET/PUT/DELETE /api/venues/{id}", "Chặn xóa nếu còn ghế hoặc suất diễn."],
            ["Performance", "GET/POST /api/performances; PUT/DELETE /api/performances/{id}", "Kiểm soát cửa sổ bán và trạng thái OnSale."],
            ["Seat area / Seat", "/api/seat-areas, /api/seats", "Section, hàng, số ghế, tọa độ và tier giá."],
            ["Ticket type / price", "/api/ticket-types; /api/performances/{id}/ticket-types", "Giá cấu hình theo từng suất diễn."],
            ["User", "/api/admin/users", "Xóa là vô hiệu hóa để giữ lịch sử."],
            ["Booking", "/api/admin/bookings", "Chỉ đổi trạng thái hợp lệ; không xóa đơn đã thanh toán tùy tiện."],
        ],
        [2000, 4600, 3000],
    )
    heading(doc, "4.3. Quy tắc nghiệp vụ", 2)
    for step in [
        "Không xóa venue khi vẫn có seat hoặc performance đang tham chiếu.",
        "Không đổi venue của performance nếu đã phát sinh seat hold hoặc ticket, vì sơ đồ ghế sẽ mất đồng nhất.",
        "Không xóa cứng user đã có booking. DELETE được xử lý theo hướng deactivate.",
        "Booking đã Paid không được hủy trực tiếp. Hoàn tiền cần quy trình riêng khi có cổng thanh toán thật.",
        "Mọi thao tác nhạy cảm đi qua authorization và audit log.",
    ]:
        body(doc, step, before=0, after=4)

    heading(doc, "5. Email xác nhận đơn vé", 1)
    heading(doc, "5.1. Mô tả chung", 2)
    body(
        doc,
        "Sau khi có booking, hệ thống xếp email xác nhận vào bảng EmailOutbox. Bản ghi lưu người nhận, tiêu đề, HTML, plain text, nhà cung cấp, số lần thử, lỗi cuối và thời điểm gửi. Đây là outbox ngoại tuyến: nội dung được kết xuất và lưu để khách xem trước, quản trị xem danh sách và gửi lại. Phiên bản hiện tại không gửi ra Gmail hoặc Outlook cho đến khi gắn bộ chuyển SMTP hoặc nhà cung cấp thư.",
    )
    heading(doc, "5.2. API", 2)
    add_table(
        doc,
        ["API", "Quyền", "Mục đích"],
        [
            ["GET /api/bookings/{id}/confirmation-email", "Đã đăng nhập", "Khách xem email của đơn mình; Staff/Admin xem được đơn được phép."],
            ["GET /api/admin/email-outbox", "Admin, Staff", "Danh sách tối đa 200 thư, lọc theo trạng thái."],
            ["GET /api/admin/email-outbox/{id}", "Admin, Staff", "Xem nội dung một thư."],
            ["POST /api/admin/email-outbox/{id}/retry", "Admin, Staff", "Đưa thư Failed về Pending rồi dispatch lại."],
            ["POST /api/admin/bookings/{id}/confirmation-email", "Admin, Staff", "Tạo thư xác nhận cho một booking."],
        ],
        [4200, 1800, 3600],
    )

    heading(doc, "6. Thiết kế dữ liệu và ràng buộc toàn vẹn", 1)
    heading(doc, "6.1. Mô hình dữ liệu chính", 2)
    add_table(
        doc,
        ["Nhóm", "Thực thể", "Thông tin chính"],
        [
            ["Tài khoản", "User", "Email, password hash, họ tên, điện thoại, role, isActive."],
            ["Catalog", "Event, Venue, Performance", "Sự kiện, địa điểm, thời gian, cửa sổ bán và trạng thái."],
            ["Sơ đồ ghế", "SeatArea, Seat, SeatInventory", "Vị trí ghế và trạng thái riêng theo từng suất diễn."],
            ["Giá vé", "TicketType, PerformanceTicketType", "Loại vé, giá cơ sở, giá và sức chứa theo suất diễn."],
            ["Giao dịch", "Booking, Ticket, Payment", "Đơn, vé phát hành và payment nội bộ."],
            ["Thông báo", "EmailOutbox", "Nội dung thư, trạng thái gửi, số lần thử và lỗi."],
            ["Vận hành", "AuditLog, SystemSetting", "Nhật ký thao tác và hash khóa scheduler."],
        ],
        [2000, 3400, 4200],
    )
    heading(doc, "6.2. Ràng buộc quan trọng", 2)
    add_table(
        doc,
        ["Ràng buộc", "Ý nghĩa", "Rủi ro được giảm"],
        [
            ["PK SeatInventory (PerformanceId, SeatId)", "Một bản ghi trạng thái cho một ghế tại một suất diễn.", "Trùng tồn kho ghế."],
            ["UNIQUE Ticket (PerformanceId, SeatId)", "Một ghế chỉ phát hành một ticket trong một suất diễn.", "Bán trùng vé do lỗi logic hoặc retry."],
            ["UNIQUE Booking (UserId, IdempotencyKey)", "Confirm gửi lại cùng key chỉ tạo một booking.", "Double-click và retry mạng."],
            ["UNIQUE Payment (BookingId, IdempotencyKey)", "Một booking không tạo payment nội bộ trùng cùng key.", "Ghi nhận giao dịch trùng."],
            ["Restrict delete", "Không xóa dữ liệu cha khi còn dữ liệu nghiệp vụ.", "Mất lịch sử và bản ghi mồ côi."],
        ],
        [3600, 3400, 2600],
    )

    heading(doc, "7. Tiến trình giải phóng hold hết hạn", 1)
    heading(doc, "7.1. Mô tả chung", 2)
    body(
        doc,
        "Hold ghế có hiệu lực 7 phút. Sau thời điểm đó, khách không còn quyền xác nhận bằng token cũ và ghế phải bán được trở lại. Đúng đắn không phụ thuộc duy nhất vào tiến trình nền: lúc đọc sơ đồ ghế và lúc giữ ghế, backend cũng nhận hold đã quá hạn. Tác vụ nền chỉ dọn chủ động để trạng thái cập nhật nhanh hơn khi không có client mở sơ đồ.",
    )
    heading(doc, "7.2. Luồng xử lý", 2)
    code_block(
        doc,
        """Scheduler (mỗi 1 phút)
  POST /api/scheduled/release-expired-holds  +  scheduler key
  Hash key SHA-256, so sánh thời gian cố định với SystemSettings
    sai hoặc thiếu key  ->  HTTP 403, không dọn
    hợp lệ              ->  tìm SeatInventory Status=Held và HoldExpiresAtUtc đã qua
                         xóa HoldToken, HeldByUserId, HoldExpiresAtUtc
                         Status -> Available
                         ghi log và trả số ghế đã giải phóng""",
    )
    body(
        doc,
        "Khóa scheduler chỉ lưu dạng hash trong database. Endpoint bảo trì từ chối request thiếu khóa hoặc sai khóa. Lịch production chạy mỗi phút và được kiểm bằng response 2xx trong nhật ký tác vụ.",
        before=8,
    )

    heading(doc, "8. Bảo mật, xử lý lỗi và quan sát", 1)
    heading(doc, "8.1. Cơ chế bảo vệ", 2)
    add_table(
        doc,
        ["Thành phần", "Cách triển khai", "Kết quả mong đợi"],
        [
            ["Authentication", "JWT trong cookie HttpOnly, SameSite=Strict; mật khẩu lưu dạng hash.", "JavaScript không đọc được token phiên."],
            ["Authorization", "Role Customer, Staff, Admin trên API.", "User thường gọi API admin nhận 403."],
            ["Input validation", "Data annotations và luật nghiệp vụ.", "Request sai trả 400, không ghi dữ liệu lỗi."],
            ["Rate limiting", "Middleware rate limiting của ASP.NET Core.", "Request bất thường bị giới hạn."],
            ["Security headers", "CSP, X-Content-Type-Options, X-Frame-Options, Referrer-Policy, Permissions-Policy.", "Giảm một số tấn công phía trình duyệt."],
            ["Error handling", "Middleware ngoại lệ toàn cục, application/problem+json.", "Không lộ stack trace; lỗi có code và traceId."],
            ["Observability", "X-Correlation-ID, request log, audit log.", "Đối chiếu request phía client với log backend."],
        ],
        [2200, 4600, 2800],
    )
    heading(doc, "8.2. Quy ước response lỗi", 2)
    body(doc, "Lỗi nghiệp vụ trả application/problem+json. Ví dụ ghế không còn bán được:")
    code_block(
        doc,
        """{
  "type": "https://httpstatuses.com/409",
  "title": "Seat is unavailable",
  "status": 409,
  "code": "seat_unavailable",
  "traceId": "..."
}""",
    )
    body(
        doc,
        "Frontend hiển thị theo code. Kỹ thuật viên dùng traceId hoặc X-Correlation-ID để rà log. API không tồn tại trả JSON 404, không để SPA fallback trả HTML 200, nên công cụ kiểm thử phân biệt được lỗi route.",
        before=8,
    )

    heading(doc, "9. Build, kiểm thử và nghiệm thu", 1)
    heading(doc, "9.1. Thành phần mã nguồn", 2)
    add_table(
        doc,
        ["Thành phần", "Thư mục / file", "Nội dung"],
        [
            ["Backend .NET", "src/NovaTickets.Api/", "Controllers, Services, Domain, Data, Infrastructure, SignalR, Swagger."],
            ["Frontend", "client/src/", "Storefront React, sơ đồ ghế, đơn vé, Admin Console và API client."],
            ["Schema", "Data/Migrations, database/", "Schema, chỉ mục, ràng buộc và script SQL 001 đến 005."],
            ["Test", "scripts/, client/src/*.test.ts", "Đồng thời ghế, luồng booking, dọn hold, bộ tích hợp."],
            ["Triển khai", "Dockerfile", "Build đa giai đoạn frontend và runtime ASP.NET Core 8."],
        ],
        [2200, 3400, 4000],
    )
    heading(doc, "9.2. Lệnh kiểm thử", 2)
    code_block(
        doc,
        """pnpm run build
pnpm test
bash scripts/test-seat-concurrency.sh
bash scripts/test-booking-flow.sh
dotnet build NovaTickets.sln""",
    )
    body(
        doc,
        "Bộ tích hợp dùng SQLite tạm ở môi trường Testing nên không đụng dữ liệu production. Phạm vi đã phủ: xung đột hai request giữ cùng ghế, nhả hold, hold hết hạn, xác nhận booking, ticket và payment, idempotency, phân quyền, validation, và retry outbox email từ Failed sang Sent.",
        before=8,
    )
    heading(doc, "9.3. Ma trận kiểm thử", 2)
    add_table(
        doc,
        ["Mã", "Test case", "Đạt", "Không đạt"],
        [
            ["TC-01", "Đăng nhập Customer", "HTTP 200, cookie JWT hợp lệ.", "Không phát cookie khi thông tin đúng."],
            ["TC-02", "Customer gọi API admin", "HTTP 403.", "Trả dữ liệu quản trị hoặc HTTP 200."],
            ["TC-03", "Giữ một ghế trống", "HTTP 200, có token và hạn; ghế Held.", "Ghế không đổi hoặc lỗi 500."],
            ["TC-04", "Hai user giữ cùng ghế", "Một 200, một 409 seat_unavailable.", "Hai 200 hoặc hai booking cho một ghế."],
            ["TC-05", "Confirm booking", "Có booking, ticket, payment; ghế Sold.", "Chỉ tạo một phần dữ liệu."],
            ["TC-06", "Retry confirm cùng key", "Trả cùng booking, không nhân đôi payment/ticket.", "Ghi giao dịch trùng."],
            ["TC-07", "Hold quá hạn", "Ghế Available, xóa token và chủ giữ.", "Ghế kẹt Held."],
            ["TC-08", "Callback scheduler", "Có key trả 200; thiếu key trả 403.", "Cho gọi không xác thực hoặc lỗi 500."],
        ],
        [1100, 2400, 3100, 3000],
    )
    heading(doc, "9.4. Tiêu chí nghiệm thu", 2)
    for step in [
        "Mọi test case bắt buộc đạt. Không còn HTTP 500 chưa được phân tích.",
        "Không có bằng chứng tạo trùng booking, ticket hoặc payment.",
        "API quản trị bị chặn với Customer. Endpoint công khai và endpoint cần quyền trả đúng mã trạng thái.",
        "Luồng đặt vé thành công tạo đủ booking, ticket, payment và đổi trạng thái ghế đúng.",
        "Tác vụ dọn hold chạy, có log 2xx, và không cho gọi khi thiếu khóa.",
    ]:
        body(doc, step, before=0, after=4)

    heading(doc, "10. Vận hành, sự cố và rollback", 1)
    heading(doc, "10.1. Xử lý sự cố", 2)
    add_table(
        doc,
        ["Sự cố", "Dấu hiệu", "Xử lý ban đầu"],
        [
            ["Không tải được sự kiện", "Frontend hiện 0 event hoặc request lỗi.", "Kiểm tra GET /api/events, traceId, log backend và kết nối database."],
            ["Ghế kẹt Held", "Hết 7 phút vẫn không chọn được.", "Xem SeatInventory, log scheduler và response 2xx của endpoint dọn."],
            ["Booking 409", "Khách không xác nhận được.", "Kiểm tra holdToken, chủ giữ, expiresAtUtc. Hướng khách chọn ghế khác nếu hold đã mất."],
            ["Booking 500", "Response có traceId.", "Tra request log và audit. Không xác nhận vé bằng tay khi chưa biết transaction đã commit hay chưa."],
            ["API admin 403", "Admin console không thao tác được.", "Kiểm tra cookie JWT, role và policy của endpoint."],
        ],
        [2400, 3200, 4000],
    )
    heading(doc, "10.2. Rollback", 2)
    body(
        doc,
        "Khi lỗi xuất hiện sau khi phát hành, ưu tiên đưa ứng dụng về bản ổn định gần nhất. Rollback ứng dụng không dùng để sửa dữ liệu giao dịch. Với migration schema, chỉ thay đổi có kiểm soát sau khi đã sao lưu và đánh giá ảnh hưởng trên production. Sau rollback, chạy lại smoke test, kiểm tra phân quyền và kiểm tra đồng thời giữ ghế.",
    )

    heading(doc, "11. Giới hạn và hướng phát triển", 1)
    add_table(
        doc,
        ["Hạng mục", "Hiện trạng", "Hướng phát triển"],
        [
            ["Payment", "Payment nội bộ dùng cho demo.", "Tích hợp VNPay, MoMo hoặc Stripe; xác minh webhook, chữ ký và đối soát."],
            ["QR / check-in", "Ticket có mã, chưa có quét tại cổng.", "QR có chữ ký, endpoint check-in một lần và màn hình nhân viên."],
            ["Hoàn tiền", "Chưa có state machine hoàn tiền thật.", "Yêu cầu hoàn, duyệt, callback nhà cung cấp và audit."],
            ["Email", "Outbox ngoại tuyến, xem trước và retry.", "Bộ chuyển SMTP hoặc nhà cung cấp thư, không đổi hợp đồng outbox."],
            ["Realtime", "SignalR phát trạng thái; frontend có polling dự phòng.", "Reconnect, số liệu kết nối và broadcast theo suất diễn."],
            ["Báo cáo", "Chưa có dashboard doanh thu chuyên sâu.", "Vé đã bán, tỉ lệ hold thành booking, doanh thu theo sự kiện và hạng ghế."],
        ],
        [1800, 3800, 4000],
    )

    heading(doc, "12. Kết luận", 1)
    body(
        doc,
        "NOVA Tickets là giải pháp web bán vé full-stack. Các việc nhạy cảm — giữ ghế, xác nhận booking, chống đặt trùng và phân quyền — nằm ở backend ASP.NET Core và cơ sở dữ liệu. Điểm then chốt là cập nhật có điều kiện, transaction, token giữ ghế, ràng buộc duy nhất và idempotency, để mỗi ghế của một suất diễn chỉ có một chủ hợp lệ. Tài liệu này ghi phạm vi, hợp đồng API, luồng xử lý, kiểm thử và vận hành cho demo, nghiệm thu và các bước phát triển sau.",
    )

    heading(doc, "Tài liệu tham khảo", 1)
    for item in [
        "NOVA Tickets, Kiến trúc hệ thống, tài liệu nội bộ dự án, 2026.",
        "NOVA Tickets, API Documentation, tài liệu nội bộ dự án, 2026.",
        "NOVA Tickets, Testing Playbook, tài liệu nội bộ dự án, 2026.",
        "Microsoft Learn, Handling Concurrency Conflicts, https://learn.microsoft.com/en-us/ef/core/saving/concurrency.",
        "Microsoft Learn, Rate limiting middleware in ASP.NET Core, https://learn.microsoft.com/en-us/aspnet/core/performance/rate-limit.",
        "OWASP, Session Management Cheat Sheet, https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html.",
    ]:
        body(doc, item, before=0, after=4)

    core = doc.core_properties
    core.title = "NOVA Tickets — Tài liệu giải pháp kỹ thuật"
    core.subject = "Giải pháp kỹ thuật hệ thống bán vé sự kiện"
    core.category = "Báo cáo"
    enable_update_fields(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
